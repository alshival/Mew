import fitz 
import markdown
import docx 
from bs4 import BeautifulSoup
import pypandoc
import openpyxl
import json
import xml.etree.ElementTree as ET

supported_filetypes = [
    '.pdf',   # PDF Documents
    '.txt',   # Plain Text Files
    '.md',    # Markdown Files
    '.docx',  # Word Documents
    '.html',  # HTML Files
    '.rtf',   # Rich Text Format Files
    '.xlsx',  # Excel Files
    '.json',  # JSON Files
    '.sh',    # Shell Script Files
    '.xml',   # XML Files
    '.csv',   # CSV Files
    '.yaml',  # YAML Files
    '.ini',   # INI Configuration Files
    '.log'    # Log Files
]

spec = {
        "type": "function",
        "function": {
            "name": "documentReader",
            "description": f"""
For extracting text from documents. Supported filetypes: {supported_filetypes}
""",
            "parameters": {
                "type": "object",
                "properties": {
                    "filepath": {
                        "type": "string",
                        "description": "Full path to the document"
                    }
                },
                "required": ["filepath"]
            }
        }
    }


def pdfReader(filepath):
    # Open the PDF file
    document = fitz.open(filepath)
    
    # Initialize an empty string to hold the text
    text = ""
    
    # Iterate through each page
    for page_num in range(len(document)):
        # Get the page
        page = document.load_page(page_num)
        # Extract the text from the page
        text += page.get_text()
    
    return text

def xmlReader(filepath):
    tree = ET.parse(filepath)
    root = tree.getroot()
    text = ET.tostring(root, encoding='unicode')
    return text


def mdToHtmlReader(filepath):
    with open(filepath, 'r', encoding='utf-8') as file:
        text = file.read()
    html = markdown.markdown(text)
    return html

def docxReader(filepath):
    document = docx.Document(filepath)
    text = ""
    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"
    return text

def htmlReader(filepath):
    with open(filepath, 'r', encoding='utf-8') as file:
        soup = BeautifulSoup(file, 'html.parser')
        text = soup.get_text()
    return text

def rtfReader(filepath):
    text = pypandoc.convert_file(filepath, 'plain')
    return text

def xlsxReader(filepath):
    workbook = openpyxl.load_workbook(filepath)
    sheet = workbook.active
    text = ""
    for row in sheet.iter_rows(values_only=True):
        text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
    return text

def jsonReader(filepath):
    with open(filepath, 'r', encoding='utf-8') as file:
        data = json.load(file)
    return json.dumps(data, indent=4)

def simpleTextReader(filepath):
    with open(filepath, 'r', encoding='utf-8') as file:
        text = file.read()
    return text

def documentReader(filepath):
    try:
        if filepath.endswith('.pdf'):
            return pdfReader(filepath)
        elif filepath.endswith(('.txt', '.md', '.sh', '.log', '.ini', '.yaml', '.yml')):
            return simpleTextReader(filepath)
        elif filepath.endswith('.docx'):
            return docxReader(filepath)
        elif filepath.endswith('.html'):
            return htmlReader(filepath)
        elif filepath.endswith('.rtf'):
            return rtfReader(filepath)
        elif filepath.endswith('.xlsx'):
            return xlsxReader(filepath)
        elif filepath.endswith('.json'):
            return jsonReader(filepath)
        elif filepath.endswith('.xml'):
            return xmlReader(filepath)
        else:
            # Attempt to read with simpleTextReader if the file type is not supported
            try:
                return simpleTextReader(filepath)
            except Exception as e:
                return f"Error: Unsupported file type or unable to read the file. Exception: {str(e)}"
    except Exception as e:
        return f"An error occurred while reading the file: {str(e)}"
